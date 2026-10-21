import csv
import io
import smtplib
import unicodedata
from datetime import date, datetime, time, timezone
from email.message import EmailMessage
from typing import Dict, List, Optional, Tuple

import httpx
import openpyxl
from icalendar import Calendar
from sqlalchemy.orm import Session

from app.config import settings
from app.models import CalendarEvent, CalendarFeed, Student
from app.schemas import ImportResult, MailSendResult


# =====================================================================
# 1. PARSER I GENERATOR ARKUSZY EXCEL (.XLSX) ORAZ PLIKÓW .CSV
# =====================================================================
COLUMN_ALIASES: Dict[str, List[str]] = {
    "imie": ["imie", "imię", "firstname", "first_name", "name", "imiona"],
    "nazwisko": ["nazwisko", "lastname", "last_name", "surname"],
    "email": ["email", "e-mail", "mail", "adres email", "adres e-mail", "poczta"],
    "grupa": ["grupa", "group", "nr grupy", "grupa dziekanska", "grupa lab", "podgrupa"],
    "komentarz": ["komentarz", "notatka", "notatki", "uwagi", "comment", "notes", "info", "status"],
}


def normalize_header(header: str) -> str:
    if not header:
        return ""
    cleaned = str(header).strip().lower()
    # Usunięcie polskich znaków diakrytycznych dla pewnego dopasowania
    nfkd = unicodedata.normalize("NFKD", cleaned)
    ascii_str = "".join(c for c in nfkd if not unicodedata.combining(c))
    return ascii_str.replace("_", " ").replace("-", " ").strip()


def map_columns(headers: List[str]) -> Dict[str, int]:
    mapped: Dict[str, int] = {}
    normalized_headers = [normalize_header(h) for h in headers]

    for field, aliases in COLUMN_ALIASES.items():
        norm_aliases = [normalize_header(a) for a in aliases]
        for idx, h in enumerate(normalized_headers):
            if h in norm_aliases or any(alias in h for alias in norm_aliases):
                if field not in mapped:
                    mapped[field] = idx
                    break
    return mapped


def extract_rows_from_file(filename: str, content: bytes) -> List[List[str]]:
    lower_name = filename.lower()
    rows: List[List[str]] = []

    if lower_name.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell).strip() for cell in row])
    elif lower_name.endswith(".csv"):
        decoded = None
        for encoding in ("utf-8-sig", "utf-8", "cp1250", "iso-8859-2"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            decoded = content.decode("utf-8", errors="replace")

        sample = decoded[:2048]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
        for row in reader:
            rows.append([cell.strip() for cell in row])
    else:
        raise ValueError("Nieobsługiwany format pliku. Dozwolone są pliki .xlsx oraz .csv.")

    return [r for r in rows if any(cell != "" for cell in r)]


def import_students_from_bytes(db: Session, filename: str, content: bytes) -> ImportResult:
    rows = extract_rows_from_file(filename, content)
    if not rows:
        return ImportResult(inserted=0, updated=0, skipped=0, errors=["Plik jest pusty."])

    headers = rows[0]
    col_map = map_columns(headers)

    required = ["imie", "nazwisko", "email"]
    missing = [req for req in required if req not in col_map]
    if missing:
        return ImportResult(
            inserted=0,
            updated=0,
            skipped=0,
            errors=[
                f"Nie znaleziono wymaganych kolumn: {', '.join(missing)}. "
                f"Wykryte nagłówki: {', '.join(headers)}"
            ],
        )

    inserted = 0
    updated = 0
    skipped = 0
    errors: List[str] = []

    for row_idx, row in enumerate(rows[1:], start=2):
        def get_val(field: str, default: str = "") -> str:
            idx = col_map.get(field)
            if idx is None or idx >= len(row):
                return default
            return row[idx].strip()

        imie = get_val("imie")
        nazwisko = get_val("nazwisko")
        email = get_val("email").lower()
        grupa = get_val("grupa", "1") or "1"
        komentarz = get_val("komentarz", "")

        if not email or "@" not in email:
            skipped += 1
            errors.append(f"Wiersz {row_idx}: Pominięto - nieprawidłowy lub pusty email ('{email}').")
            continue

        if not imie or not nazwisko:
            skipped += 1
            errors.append(f"Wiersz {row_idx}: Pominięto ({email}) - brak imienia lub nazwiska.")
            continue

        existing = db.query(Student).filter(Student.email == email).first()
        if existing:
            existing.imie = imie
            existing.nazwisko = nazwisko
            existing.grupa = grupa
            if komentarz:
                existing.komentarz = komentarz
            updated += 1
        else:
            new_student = Student(
                imie=imie,
                nazwisko=nazwisko,
                email=email,
                grupa=grupa,
                komentarz=komentarz,
            )
            db.add(new_student)
            inserted += 1

    db.commit()
    return ImportResult(inserted=inserted, updated=updated, skipped=skipped, errors=errors[:30])


def export_students_to_xlsx(students: List[Student]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Studenci RAU2"

    headers = ["ID", "Imie", "Nazwisko", "Email", "Grupa", "Komentarz", "Utworzono"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)

    for s in students:
        created_str = s.utworzono_data.strftime("%Y-%m-%d %H:%M") if s.utworzono_data else ""
        ws.append([s.id, s.imie, s.nazwisko, s.email, s.grupa, s.komentarz, created_str])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_students_to_csv(students: List[Student]) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Imie", "Nazwisko", "Email", "Grupa", "Komentarz"])
    for s in students:
        writer.writerow([s.imie, s.nazwisko, s.email, s.grupa, s.komentarz])
    return output.getvalue().encode("utf-8-sig")


# =====================================================================
# 2. KLIENT POCZTY SMTP (OFFICE 365 / OUTLOOK POLSL)
# =====================================================================
class SafeFormatDict(dict):
    def __missing__(self, key: str) -> str:
        return "{" + key + "}"


def render_mail_template(template: str, student: Student) -> str:
    data = SafeFormatDict(
        imie=student.imie,
        nazwisko=student.nazwisko,
        email=student.email,
        grupa=student.grupa,
        komentarz=student.komentarz or "",
    )
    return template.format_map(data)


def send_personalized_emails(
    students: List[Student], subject_template: str, body_template: str
) -> MailSendResult:
    if not settings.SMTP_USER or not settings.SMTP_PASSWORD:
        return MailSendResult(
            sent_count=0,
            failed_count=len(students),
            errors=[
                "Brak konfiguracji SMTP_USER lub SMTP_PASSWORD w pliku .env. "
                "Uzupełnij dane konta Office 365 lub użyj przycisku '1-Click Outlook (BCC)'."
            ],
        )

    sent_count = 0
    failed_count = 0
    errors: List[str] = []

    try:
        server = smtplib.SMTP(settings.SMTP_HOST, settings.SMTP_PORT, timeout=20)
        server.ehlo()
        if settings.SMTP_USE_TLS:
            server.starttls()
            server.ehlo()
        server.login(settings.SMTP_USER, settings.SMTP_PASSWORD)
    except Exception as exc:
        return MailSendResult(
            sent_count=0,
            failed_count=len(students),
            errors=[f"Błąd połączenia z serwerem SMTP ({settings.SMTP_HOST}:{settings.SMTP_PORT}): {exc}"],
        )

    try:
        for student in students:
            try:
                personalized_subject = render_mail_template(subject_template, student)
                personalized_body = render_mail_template(body_template, student)

                msg = EmailMessage()
                msg["Subject"] = personalized_subject
                msg["From"] = f"{settings.SMTP_FROM_NAME} <{settings.SMTP_USER}>"
                msg["To"] = f"{student.imie} {student.nazwisko} <{student.email}>"
                msg.set_content(personalized_body)

                server.send_message(msg)
                sent_count += 1
            except Exception as item_exc:
                failed_count += 1
                errors.append(f"{student.email}: {item_exc}")
    finally:
        try:
            server.quit()
        except Exception:
            pass

    return MailSendResult(sent_count=sent_count, failed_count=failed_count, errors=errors[:20])


# =====================================================================
# 3. AGREGATOR TERMINÓW: PARSER .ICS (USOSWEB & MOODLE RAU2)
# =====================================================================
CLASS_SCHEDULE_CATEGORIES = {"LAB", "WYKLAD", "LEKTORAT", "CWICZENIA", "PROJEKT"}
TIMELINE_CATEGORIES = {"KOLOKWIUM", "EGZAMIN", "DEADLINE_RAU2", "INNE"}


def classify_event_category(title: str, description: str, source_type: str) -> str:
    combined = f"{title} {description}".lower()
    if any(k in combined for k in ["egzamin", "exam"]):
        return "EGZAMIN"
    if any(k in combined for k in ["kolokwium", "wejściówka", "wejsciowka", "test", "quiz", "zaliczenie"]):
        return "KOLOKWIUM"
    if source_type == "RAU2_MOODLE" or any(
        k in combined for k in ["oddanie", "termin", "zadanie", "sprawozdanie", "due", "moodle"]
    ):
        return "DEADLINE_RAU2"
    if any(
        k in combined
        for k in ["lektorat", "język", "jezyk", "angielski", "niemiecki", "hiszpański", "spnjo", "foreign language"]
    ):
        return "LEKTORAT"
    if any(k in combined for k in ["laboratorium", "lab.", "lab ", "[lab]", "(lab)"]):
        return "LAB"
    if any(k in combined for k in ["ćwiczenia", "cwiczenia", "ćw.", "cw.", "[ćw]", "audytoryjne"]):
        return "CWICZENIA"
    if any(k in combined for k in ["projekt", "proj.", "[proj]"]):
        return "PROJEKT"
    if any(k in combined for k in ["wykład", "wyklad", "wyk.", "[wyk]", "lecture", "seminarium"]):
        return "WYKLAD"
    # Jeśli źródłem jest USOS i wydarzenie nie jest kolokwium/egzaminem, traktujemy domyślnie jako zajęcia (WYKLAD)
    if source_type == "USOS":
        return "WYKLAD"
    return "INNE"


def normalize_ical_datetime(dt_prop) -> Optional[datetime]:
    if dt_prop is None:
        return None
    val = dt_prop.dt if hasattr(dt_prop, "dt") else dt_prop
    if isinstance(val, datetime):
        if val.tzinfo is None:
            return val.replace(tzinfo=timezone.utc)
        return val.astimezone(timezone.utc)
    if isinstance(val, date):
        return datetime.combine(val, time(8, 0), tzinfo=timezone.utc)
    return None


def parse_ics_content(ics_bytes: bytes, feed: CalendarFeed) -> List[CalendarEvent]:
    cal = Calendar.from_ical(ics_bytes)
    parsed_events: List[CalendarEvent] = []

    for component in cal.walk():
        if component.name != "VEVENT":
            continue

        uid = str(component.get("UID", ""))
        summary = str(component.get("SUMMARY", "Wydarzenie bez tytułu")).strip()
        description = str(component.get("DESCRIPTION", "")).strip()
        location = str(component.get("LOCATION", "")).strip()

        start_dt = normalize_ical_datetime(component.get("DTSTART"))
        end_dt = normalize_ical_datetime(component.get("DTEND"))

        if not start_dt:
            continue

        if not uid:
            uid = f"{feed.id}-{summary}-{start_dt.isoformat()}"

        category = classify_event_category(summary, description, feed.source_type)

        parsed_events.append(
            CalendarEvent(
                feed_id=feed.id,
                uid=uid[:250],
                title=summary[:295],
                description=description,
                location=location[:250],
                start_time=start_dt,
                end_time=end_dt,
                category=category,
                source_label=feed.name,
                target_group=feed.target_group,
            )
        )

    return parsed_events


def sync_single_calendar_feed(db: Session, feed: CalendarFeed) -> Tuple[bool, int, str]:
    url = feed.ics_url.strip()
    if url.startswith("webcal://"):
        url = "https://" + url[len("webcal://") :]

    try:
        with httpx.Client(timeout=15.0, follow_redirects=True) as client:
            resp = client.get(url, headers={"User-Agent": "StarostaHub-RAU2-PolSl/1.0"})
            resp.raise_for_status()
            ics_bytes = resp.content
    except Exception as exc:
        return False, 0, f"Błąd pobierania {feed.name}: {exc}"

    try:
        new_events = parse_ics_content(ics_bytes, feed)
    except Exception as exc:
        return False, 0, f"Błąd parsowania pliku .ics ({feed.name}): {exc}"

    db.query(CalendarEvent).filter(CalendarEvent.feed_id == feed.id).delete()
    for ev in new_events:
        db.add(ev)

    feed.last_synced_at = datetime.now(timezone.utc)
    db.commit()
    return True, len(new_events), "OK"
